"""
setup_data.py
=============
One-time setup script: converts the sample CSV files into proper
Excel (.xlsx) workbooks with the correct sheet names and formatting.

Run this ONCE before running the main automation:
  python setup_data.py

Creates:
  data/Master_Tracker.xlsx       (sheet: "Master Tracker")
  data/Contractor_Update.xlsx    (sheet: "Sheet1")
"""

import sys
from pathlib import Path

# Add project root to path
PROJECT_ROOT = Path(__file__).resolve().parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

try:
    import pandas as pd
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
except ImportError:
    print("\n❌  Missing dependencies. Run first:\n  pip install pandas openpyxl\n")
    sys.exit(1)

from config import settings


def create_master_tracker() -> None:
    csv_path = PROJECT_ROOT / "data" / "Master_Tracker_Data.csv"
    xlsx_path = settings.MASTER_TRACKER_PATH

    print(f"  Creating Master Tracker from {csv_path.name}...")

    df = pd.read_csv(csv_path, dtype=str).fillna("")

    with pd.ExcelWriter(xlsx_path, engine="openpyxl") as writer:
        df.to_excel(
            writer,
            sheet_name=settings.MASTER_SHEET_NAME,
            index=False
        )

        ws = writer.sheets[settings.MASTER_SHEET_NAME]

        # Style header row
        header_fill = PatternFill("solid", fgColor="2C6E49")
        header_font = Font(bold=True, color="FFFFFF")

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        # Auto-fit column widths
        for col_idx, col in enumerate(ws.columns, 1):
            max_len = max(len(str(cell.value or "")) for cell in col)
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 4, 40)

    print(f"  [OK] Saved: {xlsx_path}")


def create_contractor_update() -> None:
    csv_path = PROJECT_ROOT / "data" / "Contractor_Update_Data.csv"
    xlsx_path = settings.CONTRACTOR_UPDATE_PATH

    print(f"  Creating Contractor Update from {csv_path.name}...")

    df = pd.read_csv(csv_path, dtype=str).fillna("")

    with pd.ExcelWriter(xlsx_path, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Sheet1", index=False)

        ws = writer.sheets["Sheet1"]

        # Style header row
        header_fill = PatternFill("solid", fgColor="1565C0")
        header_font = Font(bold=True, color="FFFFFF")

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        # Auto-fit column widths
        for col_idx, col in enumerate(ws.columns, 1):
            max_len = max(len(str(cell.value or "")) for cell in col)
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 4, 45)

    print(f"  [OK] Saved: {xlsx_path}")


def create_output_dir() -> None:
    out = settings.OUTPUT_DIR
    out.mkdir(parents=True, exist_ok=True)
    print(f"  [OK] Output directory ready: {out}")


def create_demo_inbox() -> None:
    inbox = PROJECT_ROOT / "demo_inbox"
    inbox.mkdir(parents=True, exist_ok=True)
    processed = inbox / "processed"
    processed.mkdir(parents=True, exist_ok=True)
    print(f"  [OK] Demo inbox ready: {inbox}")

    # Create a README in the demo_inbox so it's clear what to do
    readme = inbox / "README.txt"
    if not readme.exists():
        readme.write_text(
            "DROP Contractor_Update.xlsx HERE TO SIMULATE AN EMAIL ARRIVING.\n"
            "Then run: python src/main.py --mode demo\n"
            "Processed files are moved to the processed/ subfolder.\n"
        )


if __name__ == "__main__":
    print()
    print("=" * 50)
    print("  RENOVATION TRACKER -- DATA SETUP")
    print("=" * 50)
    print()

    create_master_tracker()
    create_contractor_update()
    create_output_dir()
    create_demo_inbox()

    print()
    print("  Setup complete! You can now run:")
    print()
    print("       python src/main.py")
    print("       python src/main.py --mode demo")
    print("       python src/main.py --mode gmail")
    print()
